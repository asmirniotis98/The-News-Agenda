"""Φάση 7 — αποτελέσματα αξιολόγησης από το συμπληρωμένο evaluation/eval_sheet.xlsx.

Γράφει evaluation/results.md (ποσοστά με διάστημα εμπιστοσύνης 95% Wilson, ανά όρο και ανά περίπτωση) και
evaluation/results.png (γράφημα για την αναφορά/παρουσίαση). Το results.md δεν περιέχει τίτλους ή συνόψεις —
μόνο τις κρίσεις — ώστε να μπορεί να δημοσιευτεί στο GitHub.

Εκτέλεση:  python scripts/eval_score.py
"""

import math
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
SHEET = ROOT / "evaluation" / "eval_sheet.xlsx"
OUT_MD = ROOT / "evaluation" / "results.md"
OUT_PNG = ROOT / "evaluation" / "results.png"


def wilson(k: int, n: int, z: float = 1.96) -> tuple[float, float]:
    """Διάστημα εμπιστοσύνης 95% για ποσοστό k/n (Wilson) — αξιόπιστο και για μικρά δείγματα."""
    if n == 0:
        return (0.0, 0.0)
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return (max(0.0, c - h), min(1.0, c + h))


def rate(k: int, n: int) -> str:
    if n == 0:
        return "—"
    lo, hi = wilson(k, n)
    return f"**{100 * k / n:.0f}%** ({k}/{n}, 95% CI {100 * lo:.0f}–{100 * hi:.0f}%)"


def load() -> tuple[list[dict], list[dict]]:
    if not SHEET.exists():
        sys.exit(f"Λείπει το {SHEET} — τρέξτε πρώτα: python scripts/eval_build.py")
    wb = load_workbook(SHEET, read_only=True, data_only=True)
    a = [dict(id=r[0], term=r[1], month=r[2], rank=r[3], real=r[9], event=r[10], summary=r[11], extra=r[12], notes=r[13])
         for r in wb["Α. Κορυφές"].iter_rows(min_row=2, values_only=True) if r[0]]
    b = [dict(id=r[0], term=r[1], date=r[2], relevant=r[7], notes=r[8])
         for r in wb["Β. Ακρίβεια αναζήτησης"].iter_rows(min_row=2, values_only=True) if r[0]]
    return a, b


def main() -> None:
    a, b = load()
    pending_a = [c["id"] for c in a if not (c["real"] and c["summary"] and c["extra"])]
    pending_b = [c["id"] for c in b if not c["relevant"]]
    if pending_a or pending_b:
        print(f"Προσοχή: χωρίς πλήρη βαθμολογία: {', '.join(pending_a + pending_b)} — μετρούν μόνο οι συμπληρωμένες.")

    real = Counter(c["real"] for c in a if c["real"])
    summ = Counter(c["summary"] for c in a if c["summary"])
    extra = Counter(c["extra"] for c in a if c["extra"])
    rel = Counter(c["relevant"] for c in b if c["relevant"])
    n_real = real["Ναι"] + real["Όχι"]          # τα «Αβέβαιο» αναφέρονται χωριστά
    n_summ = sum(summ.values())
    n_extra = sum(extra.values())
    n_rel = rel["Ναι"] + rel["Όχι"]

    lines = [
        "# Αξιολόγηση σε δείγμα (Φάση 7)",
        "",
        f"Παράγεται από `scripts/eval_score.py` ({date.today():%d/%m/%Y}) από το φύλλο που βαθμολόγησε άνθρωπος-αξιολογητής "
        "(`scripts/eval_build.py`). Ποσοστά με διάστημα εμπιστοσύνης 95% (Wilson) — με 30 περιπτώσεις η αβεβαιότητα είναι "
        "σημαντική και αναφέρεται ρητά.",
        "",
        "## Α. Κορυφές και συνόψεις (10 όροι × 3 ισχυρότερες κορυφές)",
        "",
        "| Μέτρο | Αποτέλεσμα |",
        "|---|---|",
        f"| Κορυφές που αντιστοιχούν σε πραγματικό γεγονός | {rate(real['Ναι'], n_real)}"
        + (f" · Αβέβαιο: {real['Αβέβαιο']}" if real["Αβέβαιο"] else "") + " |",
        f"| Συνόψεις **σωστές** | {rate(summ['Σωστή'], n_summ)} |",
        f"| Συνόψεις σωστές ή μερικώς σωστές | {rate(summ['Σωστή'] + summ['Μερικώς'], n_summ)} |",
        f"| Συνόψεις **λάθος** | {rate(summ['Λάθος'], n_summ)} |",
        f"| Συνόψεις με στοιχεία **εκτός τίτλων** («επινόηση») | {rate(extra['Ναι'], n_extra)} |",
        "",
        "### Ανά όρο",
        "",
        "| Όρος | Πραγματικές κορυφές | Σωστές συνόψεις | Εκτός τίτλων |",
        "|---|---|---|---|",
    ]
    by_term = defaultdict(list)
    for c in a:
        by_term[c["term"]].append(c)
    for term, cs in by_term.items():
        rr = sum(c["real"] == "Ναι" for c in cs); nr = sum(c["real"] in ("Ναι", "Όχι") for c in cs)
        ss = sum(c["summary"] == "Σωστή" for c in cs); ns = sum(bool(c["summary"]) for c in cs)
        ee = sum(c["extra"] == "Ναι" for c in cs); ne = sum(bool(c["extra"]) for c in cs)
        lines.append(f"| {term} | {rr}/{nr} | {ss}/{ns} | {ee}/{ne} |")

    lines += ["", "### Ανά περίπτωση", "", "| ID | Όρος | Μήνας | # | Γεγονός; | Ποιο | Σύνοψη | Εκτός τίτλων | Σημειώσεις |",
              "|---|---|---|---|---|---|---|---|---|"]
    for c in a:
        lines.append(f"| {c['id']} | {c['term']} | {c['month']} | {c['rank']} | {c['real'] or '…'} | {c['event'] or ''} | "
                     f"{c['summary'] or '…'} | {c['extra'] or '…'} | {c['notes'] or ''} |")

    lines += ["", "## Β. Ακρίβεια αναζήτησης (3 όροι × 10 τυχαία άρθρα)", "",
              "Ακρίβεια (precision) = ποσοστό των άρθρων που βρίσκει η αναζήτηση και αφορούν πράγματι τον όρο με τη σημασία "
              "που εννοούμε. Δοκιμάστηκαν όροι με γνωστούς κινδύνους: «μνημόνιο συνεργασίας», «πολιτικός σεισμός», θέμα «ακριβε…».",
              "", "| Όρος | Ακρίβεια |", "|---|---|"]
    by_term_b = defaultdict(list)
    for c in b:
        by_term_b[c["term"]].append(c)
    for term, cs in by_term_b.items():
        k = sum(c["relevant"] == "Ναι" for c in cs); n = sum(c["relevant"] in ("Ναι", "Όχι") for c in cs)
        lines.append(f"| {term} | {rate(k, n)} |")
    lines.append(f"| **Σύνολο** | {rate(rel['Ναι'], n_rel)}" + (f" · Αβέβαιο: {rel['Αβέβαιο']}" if rel["Αβέβαιο"] else "") + " |")
    misses = [c for c in b if c["relevant"] == "Όχι"]
    if misses:
        lines += ["", "Άσχετα άρθρα (για τη συζήτηση των περιορισμών):", ""]
        lines += [f"- {c['id']} ({c['term']}, {c['date']}): {c['notes'] or '—'}" for c in misses]
    if pending_a or pending_b:
        lines += ["", f"> Χωρίς πλήρη βαθμολογία: {', '.join(pending_a + pending_b)}."]
    OUT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")

    # Γράφημα: τα κύρια ποσοστά με διαστήματα εμπιστοσύνης
    metrics = [("Πραγματικές\nκορυφές", real["Ναι"], n_real), ("Σωστές\nσυνόψεις", summ["Σωστή"], n_summ),
               ("Σωστές ή\nμερικώς", summ["Σωστή"] + summ["Μερικώς"], n_summ),
               ("Χωρίς στοιχεία\nεκτός τίτλων", extra["Όχι"], n_extra), ("Ακρίβεια\nαναζήτησης", rel["Ναι"], n_rel)]
    metrics = [m for m in metrics if m[2]]
    if metrics:
        fig, ax = plt.subplots(figsize=(8, 4))
        vals = [100 * k / n for _, k, n in metrics]
        errs = [[v - 100 * wilson(k, n)[0] for v, (_, k, n) in zip(vals, metrics)],
                [100 * wilson(k, n)[1] - v for v, (_, k, n) in zip(vals, metrics)]]
        ax.bar([m[0] for m in metrics], vals, color="#2563eb", yerr=errs, capsize=5, ecolor="#1e3a5f")
        for i, (v, (_, k, n)) in enumerate(zip(vals, metrics)):
            ax.text(i, 3, f"{v:.0f}%\n({k}/{n})", ha="center", color="white", fontsize=9, fontweight="bold")
        ax.set_ylim(0, 105); ax.set_ylabel("%"); ax.set_title("Αξιολόγηση σε δείγμα (διαστήματα εμπιστοσύνης 95%)")
        ax.spines[["top", "right"]].set_visible(False)
        fig.tight_layout(); fig.savefig(OUT_PNG, dpi=150)
    print(f"✓ {OUT_MD}\n✓ {OUT_PNG}")


if __name__ == "__main__":
    main()
