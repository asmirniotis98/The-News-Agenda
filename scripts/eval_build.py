"""Φάση 7 — δημιουργία του φύλλου αξιολόγησης (evaluation/eval_sheet.xlsx).

Α. Κορυφές & συνόψεις: 10 όροι × 3 ισχυρότερες κορυφές = 30 περιπτώσεις, με τους τίτλους και τη σύνοψη του LLM.
Β. Ακρίβεια αναζήτησης: 3 «δύσκολοι» όροι × 10 τυχαία άρθρα (σταθερό seed), με απόσπασμα γύρω από τον όρο.

Οι στήλες βαθμολόγησης (κίτρινες) συμπληρώνονται από τον άνθρωπο-αξιολογητή· μετά: python scripts/eval_score.py

Οι συνόψεις χρειάζονται ANTHROPIC_API_KEY στο .env (ίδιο μοντέλο με το Space: LLM_MODEL). Χωρίς κλειδί το φύλλο
φτιάχνεται χωρίς συνόψεις και ξαναφτιάχνεται αργότερα· οι συνόψεις αποθηκεύονται (cache), οπότε πληρώνονται μία φορά.

Εκτέλεση:  python scripts/eval_build.py
"""

import random
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src.env import load_env  # noqa: E402

load_env()

import duckdb  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from src import peaks  # noqa: E402
from src.normalize import normalize  # noqa: E402
from src.search import NewsIndex, make_query  # noqa: E402
from src.summarize import MODEL, Summarizer, SummaryUnavailable, month_label  # noqa: E402

OUT = ROOT / "evaluation" / "eval_sheet.xlsx"
RAW = ROOT / "data" / "Huggingface data source" / "amna-press.parquet"

# Α: όροι από διαφορετικά πεδία (οικονομία, υγεία, μετανάστευση, φυσικές καταστροφές, διεθνή, πολιτική)
PEAK_TERMS = ["μνημόνιο", "πανδημία", "ακρίβεια", "μεταναστευτικό", "σεισμός",
              "Ουκρανία", "πυρκαγιά", "Brexit", "εκλογές", "ενεργειακή κρίση"]
PEAKS_PER_TERM = 3
TITLES_SHOWN = 15

# Β: όροι που δοκιμάζουν αδύναμα σημεία: «μνημόνιο συνεργασίας», «πολιτικός σεισμός», θέμα «ακριβε…»
PRECISION_TERMS = ["μνημόνιο", "σεισμός", "ακρίβεια"]
ARTICLES_PER_TERM = 10
SEED = 2026

GRADE = PatternFill("solid", fgColor="FFF4C2")   # κίτρινο = συμπληρώνει ο αξιολογητής
HEAD = PatternFill("solid", fgColor="1E3A5F")
WRAP = Alignment(wrap_text=True, vertical="top")


def peak_cases(idx: NewsIndex, summ: Summarizer) -> list[dict]:
    cases = []
    for term in PEAK_TERMS:
        q = make_query(term)
        df = idx.monthly(q, "text")
        pk = peaks.detect(df, "prominence", PEAKS_PER_TERM)
        for r in pk.sort_values("rank").itertuples():
            month = r.month.strftime("%Y-%m")
            titles, _ = summ.titles(q, month, "text", [])
            try:
                bullets = summ.summary(q, month, "text", [])["bullets"] if summ.enabled else []
                note = "" if summ.enabled else "χωρίς κλειδί API"
            except SummaryUnavailable as e:
                bullets, note = [], str(e)
            cases.append({"term": term, "label": q.label, "month": month, "rank": int(r.rank), "hits": int(r.hits),
                          "share": round(float(r.share), 2), "titles": titles[:TITLES_SHOWN], "n_titles": len(titles),
                          "bullets": bullets, "note": note})
            print(f"  Α {term:18} {month} #{r.rank}  {'σύνοψη ✓' if bullets else note or '—'}", flush=True)
    return cases


def precision_cases(idx: NewsIndex) -> list[dict]:
    rng = random.Random(SEED)
    raw = duckdb.connect()
    out = []
    for term in PRECISION_TERMS:
        q = make_query(term)
        match_sql, params = idx._matches(idx.con.cursor(), q, "text")
        rows = idx.con.execute(f"""
            SELECT a.id, a.created, a.category, a.title, a.url FROM articles a
            WHERE a.aid IN ({match_sql}) ORDER BY a.aid""", params).fetchall()
        sample = rng.sample(rows, min(ARTICLES_PER_TERM, len(rows)))
        texts = {}
        if RAW.exists():
            ids = [r[0] for r in sample]
            texts = dict(raw.execute(f"SELECT id, text_clean FROM read_parquet('{RAW.as_posix()}') WHERE id IN (SELECT unnest($ids))",
                                     {"ids": ids}).fetchall())
        stem = q.words[0][0]
        for aid, created, category, title, url in sorted(sample, key=lambda r: r[1]):
            out.append({"term": term, "label": q.label, "date": created.strftime("%Y-%m-%d"), "category": category,
                        "title": title, "url": url, "snippet": snippet(title + " — " + (texts.get(aid) or ""), stem)})
        print(f"  Β {term:18} {len(sample)} άρθρα από {len(rows):,}", flush=True)
    return out


def snippet(text: str, stem: str, width: int = 160) -> str:
    """Απόσπασμα γύρω από την πρώτη εμφάνιση του θέματος (η κανονικοποίηση κρατά τις θέσεις των ελληνικών γραμμάτων)."""
    text = " ".join(text.split())
    norm = normalize(text)
    i = norm.find(stem) if len(norm) == len(text) else -1
    if i < 0:
        return text[: 2 * width] + ("…" if len(text) > 2 * width else "")
    a, b = max(0, i - width), min(len(text), i + len(stem) + width)
    return ("…" if a else "") + text[a:b] + ("…" if b < len(text) else "")


def header(ws, cols: list[tuple[str, int, bool]]):
    for c, (name, width, grade) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 45


def dropdown(ws, col: str, options: str, rows: int):
    dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{rows + 1}")


def write(cases_a: list[dict], cases_b: list[dict]) -> None:
    wb = Workbook()

    info = wb.active
    info.title = "Οδηγίες"
    lines = [
        ("Φύλλο αξιολόγησης — Η Ατζέντα των Ειδήσεων (Φάση 7)", True),
        ("Συμπληρώστε μόνο τις ΚΙΤΡΙΝΕΣ στήλες (λίστα επιλογών). Μετά: python scripts/eval_score.py", False),
        ("", False),
        ("Φύλλο «Α. Κορυφές» — για κάθε κορύφωση:", True),
        ("Α1 Πραγματικό γεγονός;  Ναι = ο μήνας αντιστοιχεί σε γνωστό γεγονός/εξέλιξη για τον όρο (ελέγξτε π.χ. Wikipedia)· "
         "Όχι = δεν φαίνεται κάτι ιδιαίτερο· Αβέβαιο.", False),
        ("Α2 Ποιο γεγονός: λίγες λέξεις (π.χ. «3ο μνημόνιο»).", False),
        ("Α3 Σύνοψη:  Σωστή = αποδίδει πιστά τα κύρια θέματα των τίτλων· Μερικώς = σωστή αλλά παραλείπει το κύριο ή έχει "
         "μικρή ανακρίβεια· Λάθος = παραπλανητική ή άσχετη.", False),
        ("Α4 Κάτι εκτός τίτλων;  Ναι = η σύνοψη αναφέρει γεγονός/αριθμό/όνομα που ΔΕΝ προκύπτει από τους τίτλους "
         "(«επινόηση»)· Όχι.", False),
        ("Οι τίτλοι που εμφανίζονται είναι οι πρώτοι 15 από όσους πήρε το μοντέλο (στήλη «Τίτλοι προς LLM»).", False),
        ("", False),
        ("Φύλλο «Β. Ακρίβεια αναζήτησης» — για κάθε άρθρο:", True),
        ("Β1 Σχετικό;  Ναι = το άρθρο αναφέρεται στον όρο με τη σημασία που εννοούμε (π.χ. «μνημόνιο» = δανειακή σύμβαση· "
         "«σεισμός» = φυσικός σεισμός· «ακρίβεια» = άνοδος τιμών)· Όχι = άλλη σημασία ή άλλη λέξη· Αβέβαιο.", False),
        ("", False),
        (f"Συνόψεις: μοντέλο {MODEL}, ίδιες ρυθμίσεις με την εφαρμογή. Δείγμα Β: τυχαίο με seed {SEED} (αναπαραγώγιμο).", False),
        ("Δεδομένα: © ΑΠΕ-ΜΠΕ, glossAPI/amna-press, CC BY-NC-ND 4.0 — μόνο για την αξιολόγηση του project.", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        cell = info.cell(row=i, column=1, value=text)
        cell.font = Font(bold=bold, size=13 if i == 1 else 11)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    info.column_dimensions["A"].width = 130

    a = wb.create_sheet("Α. Κορυφές")
    header(a, [("ID", 6, False), ("Όρος", 16, False), ("Μήνας", 16, False), ("Θέση", 6, False), ("Άρθρα", 8, False),
               ("% μήνα", 8, False), ("Τίτλοι (πρώτοι 15)", 70, False), ("Τίτλοι προς LLM", 9, False),
               ("Σύνοψη LLM", 60, False), ("Α1 Πραγματικό γεγονός;", 12, True), ("Α2 Ποιο γεγονός", 26, True),
               ("Α3 Σύνοψη", 12, True), ("Α4 Κάτι εκτός τίτλων;", 12, True), ("Σημειώσεις", 30, True)])
    for i, c in enumerate(cases_a, 2):
        values = [f"A{i - 1:02d}", c["term"], month_label(c["month"]), c["rank"], c["hits"], c["share"],
                  "\n".join(f"• {t}" for t in c["titles"]), c["n_titles"],
                  "\n".join(f"• {b}" for b in c["bullets"]) or f"({c['note']})", None, None, None, None, None]
        for col, v in enumerate(values, 1):
            cell = a.cell(row=i, column=col, value=v)
            cell.alignment = WRAP
            if col >= 10:
                cell.fill = GRADE
        a.row_dimensions[i].height = 230
    dropdown(a, "J", "Ναι,Όχι,Αβέβαιο", len(cases_a))
    dropdown(a, "L", "Σωστή,Μερικώς,Λάθος", len(cases_a))
    dropdown(a, "M", "Ναι,Όχι", len(cases_a))

    b = wb.create_sheet("Β. Ακρίβεια αναζήτησης")
    header(b, [("ID", 6, False), ("Όρος", 14, False), ("Ημερομηνία", 12, False), ("Κατηγορία", 12, False),
               ("Τίτλος", 45, False), ("Απόσπασμα γύρω από τον όρο", 80, False), ("Σύνδεσμος", 18, False),
               ("Β1 Σχετικό;", 12, True), ("Σημειώσεις", 30, True)])
    for i, c in enumerate(cases_b, 2):
        values = [f"B{i - 1:02d}", c["term"], c["date"], c["category"], c["title"], c["snippet"], c["url"], None, None]
        for col, v in enumerate(values, 1):
            cell = b.cell(row=i, column=col, value=v)
            cell.alignment = WRAP
            if col >= 8:
                cell.fill = GRADE
        if c["url"]:
            b.cell(row=i, column=7).hyperlink = c["url"]
        b.row_dimensions[i].height = 110
    dropdown(b, "H", "Ναι,Όχι,Αβέβαιο", len(cases_b))

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)


def main() -> None:
    idx = NewsIndex()
    summ = Summarizer(idx)
    print(f"Συνόψεις: {'ενεργές, ' + MODEL if summ.enabled else 'ΧΩΡΙΣ κλειδί API — το φύλλο φτιάχνεται χωρίς συνόψεις'}")
    cases_a = peak_cases(idx, summ)
    cases_b = precision_cases(idx)
    if OUT.exists():
        try:
            OUT.open("r+b").close()
        except PermissionError:
            sys.exit(f"Το {OUT.name} είναι ανοιχτό στο Excel — κλείστε το και ξανατρέξτε.")
    write(cases_a, cases_b)
    print(f"✓ {OUT}  ({len(cases_a)} κορυφές, {len(cases_b)} άρθρα)")


if __name__ == "__main__":
    main()
